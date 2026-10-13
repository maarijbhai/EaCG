#!/usr/bin/env python3
"""Every statistic of sets A (clean images) and B (ideal photographs), as one Excel workbook.

For looking, not for the thesis figures (make_report_figures.py makes those):
every export in python/test_set.csv, every lead, with the per-photo, per-sheet,
per-lead-name, tracer-pair, timing, memory, classifier and photo-quality views,
and a README sheet saying exactly how each number is computed.

INPUTS  results/device/sheet_scores.csv (score_device_sheets.py), python/test_set.csv,
        each export's own header (timings, memory, diagnosis, pair_id, photo_rotation),
        android/app/src/main/assets/ml/clf_kit_labels.txt (PTB-XL labels of the sheets),
        data/device_exports/Set B/_excluded/README.txt, and the Set B photographs.
OUTPUT  results/report/stats_sets_A_B.xlsx + .json

Set A = set "clean" in the manifest (the 12 printed-sheet images, no camera);
Set B = set "ideal" (24 photographs, 2 per sheet). One extra column is
computed here and nowhere else: `rmse_mv_median_removed`, RMSE with each
signal's median removed, as SNR already is (score_lead's RMSE is not; see README).

    python make_stats_workbook.py        (from python/)
"""

from __future__ import annotations

import json
import platform
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

import numpy as np
import pandas as pd

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
sys.path.insert(0, str(HERE))

import make_report_figures as R  # noqa: E402
import run_photo as rp  # noqa: E402
import score_device_sheets as SDS  # noqa: E402

OUT = REPO / "results" / "report" / "stats_sets_A_B"
SET_NAME = {"clean": "A (clean images)", "ideal": "B (ideal photos)"}
LEAD_ORDER = ["I", "II", "III", "aVR", "aVL", "aVF", "V1", "V2", "V3", "V4", "V5", "V6", "II rhythm"]

METHODS = [
    ("WHAT THIS WORKBOOK IS", ""),
    ("Set A", "The 12 printed test sheets as clean digital images (3200 x 2304 px), picked from the phone's gallery: "
              "the app with no camera in the way. 12 images x 2 tracers."),
    ("Set B", "The same 12 sheets printed and photographed with the app in good indoor light ('ideal'): 24 photographs "
              "(2 per sheet) x 2 tracers, each pair from ONE Digitise tap in both-tracers mode (pair_id)."),
    ("Tracers", "learned = b3, the CNN (LiteRT FP16) that reads each row; classical = the threshold-and-centroid "
                "control. Everything else in the pipeline is identical for both."),
    ("", ""),
    ("WHAT THE PHONE WRITES (each CSV)", ""),
    ("Header lines", "'#' lines: device, photo, target sheet, tracer, corners, geometry (grid fit), calibration "
                     "(px/mm, ticks found of 9, time-zero source), per-stage times (ms), per-stage peak memory, "
                     "quality verdict and warnings, the classifier's probabilities, pair_id, photo_rotation."),
    ("Time series", "time_s + 13 columns: I, II, III, aVR, aVL, aVF, V1-V6 and II_rhythm, at 500 Hz, 5,000 rows "
                    "(10 s), in mV. A 3x4 lead has values only in its own 2.5 s cell (1,250 samples: I/II/III 0-2.5 s, "
                    "aVR/aVL/aVF 2.5-5 s, V1-V3 5-7.5 s, V4-V6 7.5-10 s); empty elsewhere. II_rhythm has all 10 s. "
                    "Empty = NaN = the tracer saw no trace there."),
    ("How the phone made them", "Each traced row (page pixels) is cut on the row's time knots (the printed "
                                "separator ticks), each pixel column mapped to seconds, then resampled to the 500 Hz "
                                "grid by linear interpolation (np.interp / SheetLayout.cellSignal). Amplitude: "
                                "mV = -(row - median row of the cell) / (px per mm x 10 mm/mV); so every cell is "
                                "about its own median (a printout has no absolute baseline)."),
    ("", ""),
    ("WHAT THE SCORER DOES (score_device_sheets.py -> score_12lead.score_lead)", ""),
    ("Truth", "The PTB-XL record's own 500 Hz signal (the one the kit printed), cut to exactly the same window: "
              "for a 3x4 lead the 1,250 samples the kit drew in that cell; for the rhythm strip all 5,000. No "
              "resampling of the truth: both are 500 Hz, sample k of the cell against sample k of the truth."),
    ("Gaps", "Samples the tracer left empty are filled by linear interpolation between their neighbours "
             "(run_photo.fill_nan) before every metric except coverage. A lead with < 50 finite samples is not "
             "scored at all: it appears with an empty PCC and counts as a FAILURE."),
    ("coverage", "Fraction of the cell's samples that are finite (before filling)."),
    ("pcc_asis", "Pearson correlation of truth and digitised over the cell, as cut (no shift, no rescale): "
                 "np.corrcoef. Offset- and scale-invariant: it measures SHAPE AND TIMING, not amplitude."),
    ("snr_db_asis", "10 log10( sum truth^2 / sum (truth - digitised)^2 ), both with their own median removed. "
                    "Sensitive to shape, timing AND amplitude."),
    ("shifted_snr_db_asis", "The best SNR over integer shifts of -25..+25 samples (+-50 ms at 500 Hz), each scored "
                            "on the overlapping part only: the PhysioNet Challenge 2024 convention, which forgives a "
                            "small time offset."),
    ("rmse_mv_asis", "sqrt(mean (truth - digitised)^2) in mV, AS THE SCORER COMPUTES IT: truth in its raw mV (its "
                     "own baseline offset kept), digitised about its median. NOTE: this counts the truth window's "
                     "baseline offset (median ~0.035 mV, up to 0.155) as error, which no digitiser could remove; "
                     "SNR and the phone's own RMSE remove both medians. See rmse_mv_median_removed."),
    ("rmse_mv_median_removed", "Computed in THIS workbook only: the same RMSE with both medians removed (consistent "
                               "with SNR). The fair amplitude error."),
    ("gain_ratio", "Least-squares slope of digitised on truth: cov(truth, dig) / var(truth). 1.0 = amplitude right; "
                   "0.9 = 10% low. Intercept fitted, so a baseline error does not affect it."),
    ("r_ratio", "R-peak height ratio: truth peaks = local maxima above half the largest deflection (polarity taken "
                "from the truth, so a downward QRS is scored on its S wave), >= 0.25 s apart; heights from each "
                "signal's median; for each peak the digitised value is its largest same-polarity deviation within "
                "+-25 ms; the median of digitised / truth over the peaks."),
    ("pcc_aligned / align_*", "The best PCC after letting the digitised signal be rescaled in time (x0.90 to x1.10 "
                              "in 0.25% steps, np.interp) and shifted (+-100 ms for a 2.5 s cell; any lag by FFT "
                              "cross-correlation for the 10 s strip). It shows what the trace would score with a perfect "
                              "time base: high aligned PCC with lower plain PCC = a TIMING error, not a reading error. "
                              "Withheld (empty, 'n.m.') when the aligned overlap is < 85% of the window."),
    ("shift_ms / shift_status", "The integer lag (+-200 ms) maximising PCC of truth against digitised, medians removed, "
                                "NaN pairs skipped, >= 85% overlap; status UNRELIABLE if that PCC < 0.5, FAILED at the "
                                "search bound. A diagnostic of where the time error is."),
    ("", ""),
    ("HOW THE SUMMARIES ARE MADE", ""),
    ("Unit", "The LEAD: 12 per photo (the rhythm strip is reported separately). Medians and quartiles over leads; a "
             "lead with no score is a failure: counted in n, n_unavailable and as 'not >= 0.9', excluded from the "
             "median."),
    ("95% CI", "Of a median: CLUSTER bootstrap over photos (leads of one photo are not independent), 2,000 resamples, "
               "seed 20261012 (make_report_figures.cluster_ci)."),
    ("Learned vs classical", "Paired by (photo, lead); Wilcoxon signed-rank on the PCC difference."),
    ("Times", "Per export, from the header's timings_ms; 'pipeline' = every stage except score* (scoring against "
              "the truth is test harness, not the app). Memory: the peak native heap per stage (MB), from "
              "memory_peak_mb."),
    ("Photo quality (Set B)", "Measured on each photograph here: sheet fill (quad area / frame), corner angles' "
                              "largest departure from 90 degrees (the phone not held parallel), keystone (longer / "
                              "shorter of top and bottom edges), light ratio (5th / 95th percentile of the paper's "
                              "smoothed luma; 1 = perfectly even), sharpness (variance of the Laplacian), clipped "
                              "fraction (pixels >= 250)."),
]


def git(*a) -> str:
    return subprocess.run(["git", *a], cwd=REPO, capture_output=True, text=True).stdout.strip()


def header(path: Path) -> dict:
    h = {}
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.startswith("#"):
            break
        k, _, v = line[1:].partition(":")
        h[k.strip()] = v.strip()
    return h


def rmse_median_removed(manifest: pd.DataFrame) -> pd.DataFrame:
    rows = []
    for _, r in manifest.iterrows():
        p = SDS.find_export(r.export)
        meta, cols = SDS.read_export(p)
        for name, rhythm, _, _, sig, tr in SDS.export_leads(p, meta, cols):
            if np.isfinite(sig).sum() < 50:
                continue
            s = rp.fill_nan(sig)
            k = min(len(s), len(tr))
            a, b = tr[:k] - np.median(tr[:k]), s[:k] - np.median(s[:k])
            rows.append({"export": r.export, "lead": name, "rhythm": bool(rhythm),
                         "rmse_mv_median_removed": round(float(np.sqrt(np.mean((a - b) ** 2))), 4)})
    return pd.DataFrame(rows)


def kit_labels() -> dict:
    out = {}
    for line in (REPO / "android/app/src/main/assets/ml/clf_kit_labels.txt").read_text(encoding="utf-8").splitlines():
        if line.startswith("#") or not line.strip():
            continue
        kv = dict(p.split("=", 1) for p in line.split())
        out[int(kv["ecg_id"])] = kv["superclasses"]
    return out


def photo_quality(path: Path) -> dict:
    import cv2
    from ecgproto import page
    from run_orientation import load
    img = load(path)
    hgt, wid = img.shape[:2]
    c = page.find_sheet_corners(img).corners.reshape(4, 2)
    quad = c.astype(np.float32)
    top, bot = np.linalg.norm(c[1] - c[0]), np.linalg.norm(c[2] - c[3])
    ang = []
    for i in range(4):
        u, v = c[i - 1] - c[i], c[(i + 1) % 4] - c[i]
        ang.append(np.degrees(np.arccos(np.dot(u, v) / np.linalg.norm(u) / np.linalg.norm(v))))
    g = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)
    m = np.zeros_like(g)
    cv2.fillConvexPoly(m, quad.astype(np.int32), 1)
    er = cv2.erode(m, np.ones((41, 41), np.uint8)) > 0
    bg = cv2.resize(cv2.medianBlur(cv2.resize(g, None, fx=0.25, fy=0.25), 31), (wid, hgt))
    vals = bg[er].astype(float)
    return {"frame": f"{wid}x{hgt}", "orientation_shot": "landscape" if wid > hgt else "portrait",
            "sheet_fill": round(float(cv2.contourArea(quad) / (wid * hgt)), 3),
            "max_corner_angle_dev_deg": round(float(np.max(np.abs(np.array(ang) - 90))), 2),
            "keystone_top_bottom": round(float(max(top, bot) / min(top, bot)), 3),
            "light_ratio_p5_p95": round(float(np.percentile(vals, 5) / np.percentile(vals, 95)), 3),
            "paper_luma_median": round(float(np.median(vals)), 1),
            "sharpness_lap_var": round(float(cv2.Laplacian(g, cv2.CV_64F)[er].var()), 1),
            "clipped_fraction": round(float((g[er] >= 250).mean()), 5)}


def summarise(df: pd.DataFrame, by: list[str]) -> pd.DataFrame:
    rows = []
    for key, g in df.groupby(by, sort=False):
        key = key if isinstance(key, tuple) else (key,)
        sm = R.summary(g["pcc_asis"])
        row = dict(zip(by, key))
        row.update(photos=g["photo_id"].nunique(), leads=sm["n"], leads_unavailable=sm["n_unavailable"],
                   pct_pcc_ge_0_9=round(sm["pct_ge_09"], 1),
                   pcc_median=sm["median"], pcc_q1=sm["q1"], pcc_q3=sm["q3"])
        lo, hi = R.cluster_ci(g, "pcc_asis")
        row.update(pcc_ci95_lo=lo, pcc_ci95_hi=hi)
        for col in ("pcc_aligned", "snr_db_asis", "shifted_snr_db_asis", "rmse_mv_asis", "rmse_mv_median_removed",
                    "gain_ratio", "r_ratio", "coverage"):
            v = pd.to_numeric(g[col], errors="coerce").dropna()
            row[f"{col}_median"] = float(v.median()) if len(v) else np.nan
        row["pcc_aligned_not_measurable"] = int(pd.to_numeric(g["pcc_aligned"], errors="coerce").isna().sum()
                                                - g["pcc_asis"].isna().sum())
        rows.append(row)
    return pd.DataFrame(rows)


def main() -> int:
    manifest = pd.read_csv(R.MANIFEST)
    manifest = manifest[manifest["set"].isin(SET_NAME)]
    df, _ = R.load(R.SCORES, R.MANIFEST)          # the report's own join, types and tracer names
    df = df[df["set"].isin(SET_NAME)].copy()
    df["set_name"] = df["set"].map(SET_NAME)
    df = df.merge(rmse_median_removed(manifest), on=["export", "lead", "rhythm"], how="left")
    df["lead_label"] = np.where(df["rhythm"], "II rhythm", df["lead"])
    labels = kit_labels()
    df["ptbxl_label"] = df["ecg_id"].map(labels)
    leads12 = df[~df["rhythm"]]

    # ---- per photo/export
    per = []
    for exp, g in df.groupby("export", sort=False):
        h = header(SDS.find_export(exp))
        l12 = g[~g["rhythm"]]
        tm = R.parse_timings(h.get("timings_ms"))
        pipe = {k: v for k, v in tm.items() if R.stage_kind(k) == "pipeline"}
        base, mem = R.parse_memory(h.get("memory_peak_mb"))
        dx = dict(re.findall(r"(\w+)=([0-9.]+)", h.get("diagnosis", "")))
        thr = dict(re.findall(r"(\w+)=([0-9.]+)", h.get("diagnosis_thresholds", "")))
        flags = " ".join(k for k in dx if k in thr and float(dx[k]) >= float(thr[k]))
        row = {"set": g["set_name"].iloc[0], "ecg_id": int(g["ecg_id"].iloc[0]), "ptbxl_label": g["ptbxl_label"].iloc[0],
               "photo_id": g["photo_id"].iloc[0], "tracer": g["tracer_kind"].iloc[0], "export": exp,
               "photo": h.get("photo"), "photo_path": g["photo_path"].iloc[0], "photo_rotation": h.get("photo_rotation", ""),
               "corners": h.get("corners"), "geometry": h.get("geometry"), "calibration": h.get("calibration"),
               "quality": h.get("quality"), "warnings": h.get("warnings", ""),
               "pcc_median_12": float(l12["pcc_asis"].median()), "pcc_min_12": float(l12["pcc_asis"].min()),
               "leads_pcc_ge_0_9": int((l12["pcc_asis"] >= 0.9).sum()), "leads_unavailable": int(l12["pcc_asis"].isna().sum()),
               "pcc_aligned_median_12": float(pd.to_numeric(l12["pcc_aligned"], errors="coerce").median()),
               "snr_db_median_12": float(l12["snr_db_asis"].median()),
               "shifted_snr_db_median_12": float(l12["shifted_snr_db_asis"].median()),
               "rmse_mv_median_12": float(l12["rmse_mv_asis"].median()),
               "rmse_mv_median_removed_median_12": float(l12["rmse_mv_median_removed"].median()),
               "gain_median_12": float(l12["gain_ratio"].median()),
               "pcc_rhythm": float(g.loc[g["rhythm"], "pcc_asis"].iloc[0]) if g["rhythm"].any() else np.nan,
               "pipeline_ms": int(sum(pipe.values())) if pipe else np.nan,
               "trace_ms": pipe.get("trace", np.nan),
               "peak_native_mb": max(v[0] for v in mem.values()) if mem else np.nan,
               "pss_end_mb": h.get("pss_end_mb", ""),
               "classifier_flags": flags}
        for k, v in dx.items():
            row[f"p_{k}"] = float(v)
        per.append(row)
    per = pd.DataFrame(per).sort_values(["set", "ecg_id", "photo_id", "tracer"])

    # ---- timing & memory per stage
    tim = []
    for _, r in per.iterrows():
        h = header(SDS.find_export(r["export"]))
        tm = R.parse_timings(h.get("timings_ms"))
        _, mem = R.parse_memory(h.get("memory_peak_mb"))
        row = {"set": r["set"], "ecg_id": r["ecg_id"], "photo_id": r["photo_id"], "tracer": r["tracer"]}
        row.update({f"ms_{k}": v for k, v in tm.items()})
        row.update({f"native_mb_{k}": v[0] for k, v in mem.items()})
        tim.append(row)
    tim = pd.DataFrame(tim)

    # ---- learned vs classical, per photo and per lead
    pairs, _ = R.paired_tracers(leads12)
    pairs["diff"] = pairs["learned"] - pairs["classical"]
    pp = (pairs.groupby("photo_id").agg(leads=("diff", "size"), learned_pcc_median=("learned", "median"),
                                        classical_pcc_median=("classical", "median"), diff_median=("diff", "median"),
                                        learned_better_leads=("diff", lambda d: int((d > 0).sum())))
          .reset_index())
    pp["set"] = pp["photo_id"].map(leads12.drop_duplicates("photo_id").set_index("photo_id")["set_name"])
    w = {s: R.wilcoxon(pairs[pairs["photo_id"].isin(leads12.loc[leads12["set_name"] == s, "photo_id"])])
         for s in SET_NAME.values()}

    # ---- views
    summary = summarise(leads12, ["set_name", "tracer_kind"])
    rhythm = summarise(df[df["rhythm"]], ["set_name", "tracer_kind"])
    per_sheet = summarise(leads12, ["set_name", "ecg_id", "tracer_kind"]).sort_values(["set_name", "ecg_id", "tracer_kind"])
    per_sheet.insert(2, "ptbxl_label", per_sheet["ecg_id"].map(labels))
    per_lead = (df.groupby(["set_name", "tracer_kind", "lead_label"])["pcc_asis"].median().unstack("lead_label")
                .reindex(columns=LEAD_ORDER).reset_index())
    per_lead_snr = (df.groupby(["set_name", "tracer_kind", "lead_label"])["snr_db_asis"].median().unstack("lead_label")
                    .reindex(columns=LEAD_ORDER).reset_index())

    # ---- photo quality (Set B)
    pq = []
    for ph, g in per[per["photo_path"].fillna("").astype(str).str.len() > 0].groupby("photo_path"):
        q = photo_quality(REPO / ph)
        q.update(photo=ph, ecg_id=int(g["ecg_id"].iloc[0]),
                 px_per_mm=float(re.search(r"([\d.]+) px/mm", g["calibration"].iloc[0]).group(1)),
                 learned_pcc_median=float(g.loc[g["tracer"] == "learned", "pcc_median_12"].iloc[0]),
                 classical_pcc_median=float(g.loc[g["tracer"] == "classical", "pcc_median_12"].iloc[0]))
        pq.append(q)
        print("photo", ph, flush=True)
    pq = pd.DataFrame(pq).sort_values(["ecg_id", "photo"]) if pq else pd.DataFrame()
    if len(pq):
        from scipy.stats import spearmanr
        corr = []
        for col in ("max_corner_angle_dev_deg", "keystone_top_bottom", "sheet_fill", "px_per_mm",
                    "light_ratio_p5_p95", "sharpness_lap_var"):
            for t in ("learned", "classical"):
                r_ = spearmanr(pq[col], pq[f"{t}_pcc_median"])
                corr.append({"photo measure": col, "tracer": t, "spearman_rho": round(float(r_.statistic), 3),
                             "p": round(float(r_.pvalue), 4), "n photos": len(pq)})
        corr = pd.DataFrame(corr)
    else:
        corr = pd.DataFrame()

    excl_file = REPO / "data" / "device_exports" / "Set B" / "_excluded" / "README.txt"
    excl = (pd.DataFrame([l.split("\t", 1) for l in excl_file.read_text(encoding="utf-8").splitlines()
                          if l and not l.startswith("#")], columns=["file", "reason"]) if excl_file.exists() else pd.DataFrame())

    # ---- write
    OUT.parent.mkdir(parents=True, exist_ok=True)
    xl = OUT.with_suffix(".xlsx")
    lead_cols = ["set_name", "ecg_id", "ptbxl_label", "photo_id", "tracer_kind", "export", "lead_label", "row", "col",
                 "coverage", "pcc_asis", "pcc_aligned", "snr_db_asis", "shifted_snr_db_asis", "rmse_mv_asis",
                 "rmse_mv_median_removed", "gain_ratio", "r_ratio", "r_ratio_aligned", "align_scale_pct",
                 "align_lag_ms", "align_overlap_s", "shift_ms", "shift_status", "quality", "warnings"]
    sheets = {
        "README": pd.DataFrame(METHODS, columns=["item", "what it is / how it is computed"]),
        "Summary 12 leads": summary,
        "Summary rhythm strip": rhythm,
        "Tracer comparison": pd.DataFrame([{"set": s, **v} for s, v in w.items()]),
        "Per photo": per,
        "Per sheet": per_sheet,
        "PCC by lead": per_lead,
        "SNR by lead": per_lead_snr,
        "Learned vs classical": pp,
        "Timing & memory": tim,
        "Photo quality (B)": pq,
        "Quality vs PCC (B)": corr,
        "Per lead (all)": df[lead_cols].sort_values(["set_name", "ecg_id", "photo_id", "tracer_kind", "row", "col"]),
        "Excluded (Set B)": excl,
    }
    with pd.ExcelWriter(xl, engine="openpyxl") as wr:
        for name, d in sheets.items():
            d.to_excel(wr, sheet_name=name, index=False)
            ws = wr.sheets[name]
            ws.freeze_panes = "A2"
            for i, col in enumerate(d.columns, start=1):
                width = max([len(str(col))] + [len(str(v)) for v in d[col].head(200).tolist()])
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(max(9, width + 2),
                                                                             110 if name == "README" else 48)
                for row in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    for cell in row:
                        if isinstance(cell.value, float):
                            cell.number_format = "0.000"
            if name == "README":
                from openpyxl.styles import Alignment, Font
                for row in ws.iter_rows(min_row=2):
                    row[1].alignment = Alignment(wrap_text=True, vertical="top")
                    if row[0].value and not row[1].value:
                        row[0].font = Font(bold=True)
    OUT.with_suffix(".json").write_text(json.dumps(R.clean_json({
        "generated_by": "python/make_stats_workbook.py",
        "utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "git_commit": git("rev-parse", "HEAD"), "git_dirty": bool(git("status", "--porcelain", "--", "python")),
        "inputs": {"scores": str(R.SCORES.relative_to(REPO)), "manifest": str(R.MANIFEST.relative_to(REPO))},
        "sheets": {k: int(len(v)) for k, v in sheets.items()},
        "versions": {"python": platform.python_version(), "pandas": pd.__version__, "numpy": np.__version__},
    }), indent=2) + "\n", encoding="utf-8")
    print(f"wrote {xl}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
